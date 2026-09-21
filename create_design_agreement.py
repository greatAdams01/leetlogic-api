from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "/Users/greatadams/Documents/ChatGPT/LEETLOGIC/Leetlogic_Design_Services_Agreement_David_Chikwendu.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = OxmlElement("w:keepNext")
    p_pr.append(keep_next)


def set_cell_text(cell, text, bold=False, color="000000", size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos"
    r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    r._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_clause(doc, number, title, paragraphs=None, bullets=None):
    h = doc.add_paragraph(style="Heading 1")
    h.add_run(f"{number}. {title}")
    set_keep_with_next(h)
    for text in paragraphs or []:
        p = doc.add_paragraph(text)
        p.paragraph_format.keep_together = True
    for text in bullets or []:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text)
        p.paragraph_format.keep_together = True


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(34, 34, 34)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(24)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)

for name, size in (("Heading 1", 13), ("Heading 2", 11)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.keep_with_next = True

for list_style in ("List Bullet", "List Number"):
    styles[list_style].font.name = "Aptos"
    styles[list_style].font.size = Pt(10.5)
    styles[list_style].paragraph_format.space_after = Pt(3)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Leetlogic Design Services Agreement")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Brand Identity User Experience and User Interface Design")
r.bold = True
r.font.size = Pt(12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Effective date: ____________________")
r.italic = True
r.font.size = Pt(10)

intro = doc.add_paragraph()
intro.add_run("This Design Services Agreement (the “Agreement”) is made between ").bold = False
intro.add_run("Great Ifeanyichukwu Adams").bold = True
intro.add_run(", email: kingifean@gmail.com (the “Contractor”), and ")
intro.add_run("David Chikwendu").bold = True
intro.add_run(", email: thedavelement@gmail.com (the “Designer”). The Contractor is engaging the Designer to provide branding, user experience, and user interface design services for the Leetlogic application and related digital products (the “Project”).")

summary = doc.add_table(rows=8, cols=2)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
summary.autofit = False
summary.columns[0].width = Inches(1.75)
summary.columns[1].width = Inches(5.0)
labels = [
    ("Project", "Leetlogic branding and UI/UX design"),
    ("Designer", "David Chikwendu · thedavelement@gmail.com"),
    ("Contractor", "Great Ifeanyichukwu Adams · kingifean@gmail.com"),
    ("Professional fee", "₦400,000 (Four Hundred Thousand Naira)"),
    ("Total project budget", "₦435,000 including the ₦35,000 tools allocation"),
    ("Payment", "60% of professional fee upfront (₦240,000); balance on completion (₦160,000)"),
    ("Tools allocation", "₦35,000 reimbursed on completion with proof of purchase"),
    ("Timeline", "Maximum 3–4 weeks, including wireframes, interactive prototypes, and final designs"),
]
for i, (label, value) in enumerate(labels):
    set_cell_text(summary.cell(i, 0), label, bold=True, color="FFFFFF")
    set_cell_shading(summary.cell(i, 0), "1F6F50")
    set_cell_text(summary.cell(i, 1), value)

add_clause(doc, 1, "Appointment and purpose", [
    "The Contractor appoints the Designer, and the Designer accepts the appointment, to create the brand identity and end-to-end UI/UX design for Leetlogic in accordance with this Agreement and the approved project brief.",
    "The Designer will collaborate with the Contractor as the primary project contact. Instructions or feedback from Leetlogic or any other stakeholder will become binding on the Designer only when confirmed or relayed by the Contractor in writing."
])

add_clause(doc, 2, "Scope of services")

h = doc.add_paragraph(style="Heading 2")
h.add_run("2.1 Brand identity")
for item in [
    "Creative direction consistent with Leetlogic’s farmer-first, trustworthy, accessible, Nigerian agritech positioning.",
    "Primary logo, secondary logo or lockup where useful, monochrome and reversed versions, favicon, and square app-icon treatment.",
    "Refinement of the proposed deep green, harvest gold, off-white, charcoal, and optional trust-blue palette, including accessible colour specifications.",
    "Typography, iconography, image style, spacing, and basic visual-usage rules.",
    "A concise brand guideline document and organized export package containing appropriate editable source files and common production formats."
]:
    doc.add_paragraph(item, style="List Bullet")

h = doc.add_paragraph(style="Heading 2")
h.add_run("2.2 User experience design")
for item in [
    "Information architecture and user flows for farmer, buyer, and administrative experiences.",
    "Low-fidelity wireframes for the principal journeys, including language selection, phone/OTP onboarding, farmer verification, product listing, listings and sales, catalogue browsing, product details, checkout, order tracking, disputes, and staff review workflows.",
    "A picture-led, low-literacy design approach using large tap targets, minimal typing, clear progress indicators, forgiving navigation, and understandable status messages.",
    "Consideration of low-bandwidth use, offline listing drafts, multilingual text expansion, voice-note entry, and text-to-speech controls at the interface-design level."
]:
    doc.add_paragraph(item, style="List Bullet")

h = doc.add_paragraph(style="Heading 2")
h.add_run("2.3 User interface design")
for item in [
    "High-fidelity, responsive designs for the Android-first farmer application.",
    "High-fidelity, responsive buyer web experience, including browsing, filtering, product details, checkout, order status, profile, and help states.",
    "High-fidelity web admin dashboard covering account verification, international-listing review, disputes, flagged listings, and essential reporting views.",
    "Marketing website designs for Home, About, How It Works, For Buyers, Impact, Founder or Team, FAQ, and Contact pages.",
    "A reusable component library covering navigation, buttons, form controls, cards, tables, badges, modals, alerts, loading, empty, offline, success, and error states.",
    "A clickable prototype of the principal farmer and buyer journeys, plus an organized developer handoff in Figma or another mutually approved design platform."
]:
    doc.add_paragraph(item, style="List Bullet")

add_clause(doc, 3, "Design requirements", bullets=[
    "The primary slogan is “Sell Local. Reach Global.” The secondary trust-focused tagline is “No Middlemen. Just Farmers and You.”",
    "The work must remain legible on small screens and lower-end devices and should be evaluated for use in bright outdoor conditions.",
    "The logo must remain recognizable at favicon and approximately 48–60 pixel app-icon sizes and must work in full colour and one colour.",
    "The language selector must be visible before sign-up and remain accessible after onboarding.",
    "Design files must anticipate English, Nigerian Pidgin, and future Nigerian-language content without embedding interface text into images.",
    "Shipping cost, delivery estimates, pricing, verification indicators, payment status, and relevant trust signals must be clearly presented in the applicable designs."
])

add_clause(doc, 4, "Deliverables and handoff", paragraphs=[
    "The Designer will provide the following final deliverables in an organized project folder or shared design workspace:"
], bullets=[
    "Editable logo and brand source files, with SVG, PNG, and PDF exports where applicable.",
    "Brand guideline document covering logo use, colour, typography, icons, imagery, and common digital applications.",
    "Editable UX flows, wireframes, high-fidelity screens, prototypes, and component library.",
    "Developer-ready specifications, including dimensions, spacing, states, responsive behaviour, colours, typography, and exportable assets.",
    "A final handoff session of up to two hours to explain the files and answer implementation questions."
])

p = doc.add_paragraph()
p.add_run("Completion requirement. ").bold = True
p.add_run("The Designer must submit the final design package to the Contractor for delivery to the Leetlogic client. The Project is not complete until the client has reviewed and approved the final design in writing and the Designer has supplied all agreed final deliverables and editable files.")

add_clause(doc, 5, "Items outside the fee", paragraphs=[
    "Unless the parties add them through a written change request, the fee does not include software development, website or application coding, production deployment, paid fonts or stock assets, professional translation, copywriting beyond short interface labels, printing, animation, photography, user-research recruitment costs, regulatory advice, or third-party subscription charges."
])

add_clause(doc, 6, "Schedule and cooperation", paragraphs=[
    "The Project duration must not exceed three to four weeks from the agreed start date. This period includes brand development, user flows, wireframes, high-fidelity website and application designs, interactive prototypes, final client review, approved revisions, and delivery of the complete final design package.",
    "The start date and interim milestone dates will be agreed in writing before work begins and may be recorded by email or in an approved project schedule. The Designer will plan submissions early enough to allow the client’s review and approval within the three-to-four-week period and will promptly notify the Contractor of any expected delay.",
    "The Contractor will provide the current Leetlogic briefs, consolidated feedback, required copy, stakeholder decisions, and reasonable access to relevant project information. The Contractor will use reasonable efforts to obtain and return client feedback promptly. A delay caused by materials, access, or feedback not being supplied when reasonably requested will extend the affected deadline by the period of that delay, provided the Designer gives prompt written notice."
])

add_clause(doc, 7, "Fees and payment")
budget = doc.add_table(rows=6, cols=2)
budget.alignment = WD_TABLE_ALIGNMENT.CENTER
budget.autofit = False
budget.columns[0].width = Inches(4.9)
budget.columns[1].width = Inches(1.8)
for j, text in enumerate(("Deliverable", "Amount")):
    set_cell_text(budget.cell(0, j), text, bold=True, color="FFFFFF")
    set_cell_shading(budget.cell(0, j), "1F6F50")
set_repeat_table_header(budget.rows[0])
budget_rows = [
    ("Branding", "₦50,000"),
    ("Website design", "₦150,000"),
    ("App design", "₦200,000"),
    ("Figma Pro or AI tools", "₦35,000"),
    ("Total project budget", "₦435,000"),
]
for i, row in enumerate(budget_rows, 1):
    for j, text in enumerate(row):
        set_cell_text(budget.cell(i, j), text, bold=(i == len(budget_rows)))
        if i % 2 == 0 or i == len(budget_rows):
            set_cell_shading(budget.cell(i, j), "F3F7F5")

h = doc.add_paragraph(style="Heading 2")
h.add_run("7.1 Payment schedule")
fees = doc.add_table(rows=4, cols=4)
fees.alignment = WD_TABLE_ALIGNMENT.CENTER
fees.autofit = False
widths = [1.85, 1.0, 1.25, 2.45]
headers = ["Payment", "Share", "Amount", "Due"]
for j, text in enumerate(headers):
    fees.columns[j].width = Inches(widths[j])
    set_cell_text(fees.cell(0, j), text, bold=True, color="FFFFFF")
    set_cell_shading(fees.cell(0, j), "1F6F50")
set_repeat_table_header(fees.rows[0])
rows = [
    ("Initial fee payment", "60%", "₦240,000", "Before the Designer begins work"),
    ("Final fee payment", "40%", "₦160,000", "After client approval and project completion"),
    ("Tools reimbursement", "—", "₦35,000", "On completion, with proof of purchase"),
]
for i, row in enumerate(rows, 1):
    for j, text in enumerate(row):
        set_cell_text(fees.cell(i, j), text)
        if i % 2 == 0:
            set_cell_shading(fees.cell(i, j), "F3F7F5")

for text in [
    "The professional design fee is ₦400,000 (Four Hundred Thousand Naira). The separate ₦35,000 tools allocation brings the total project budget to ₦435,000 (Four Hundred and Thirty-Five Thousand Naira).",
    "The initial payment is an advance against the professional-fee deliverables and will be credited when calculating amounts earned. If the Agreement ends before completion, it is subject to the reconciliation in Section 13; only the value of completed, client-approved deliverables is retained as earned professional fees.",
    "Any bank, transfer, or platform charges imposed on the recipient’s side will be borne by the Designer unless the parties agree otherwise in writing.",
    "Final editable and production-ready files may be withheld until all amounts due under this Agreement have been paid."
]:
    doc.add_paragraph(text)

h = doc.add_paragraph(style="Heading 2")
h.add_run("7.2 Design tools allocation")
for text in [
    "In addition to the ₦400,000 professional fee, ₦35,000 is allocated for one month of Figma Pro or another appropriate premium AI design tool used specifically to improve the Designer’s efficiency and delivery speed on the Project.",
    "The Designer will initially pay for the approved subscription. The subscription cost is a non-refundable Project expense once purchased for commencement of the work. The Designer must provide a valid receipt, invoice, or other reasonable proof of purchase.",
    "The ₦35,000 will be reimbursed upon completion of the Project together with the final fee payment. If this Agreement is terminated after the approved subscription has been purchased but before Project completion, the documented tools allocation will become payable as part of the termination reconciliation. It is separate from the Designer’s professional fee.",
    "The Designer is responsible for cancelling or renewing the subscription after the reimbursed one-month period. Any further subscription cost requires separate written approval from the Contractor."
]:
    doc.add_paragraph(text)

add_clause(doc, 8, "Review revisions and acceptance", paragraphs=[
    "The fee includes up to two consolidated revision rounds for the brand direction and up to two consolidated revision rounds for each approved major UI/UX stage: wireframes and high-fidelity design. A revision round means one combined list of comments supplied by the Contractor after reviewing the relevant submission.",
    "New features, a new creative direction after approval, repeated changes caused by conflicting stakeholder feedback, or work outside Section 2 will be treated as additional work. The Designer must obtain the Contractor’s written approval of the additional fee and schedule before beginning it.",
    "The Designer will submit the final design through the Contractor for review and written approval by the Leetlogic client. A submission, prototype demonstration, or Contractor review alone does not constitute final completion. The Project will be considered complete only when the client has approved the final design in writing, the Designer has addressed the approved final feedback within the agreed scope, and all deliverables in Section 4 have been supplied.",
    "The Contractor will consolidate client feedback and communicate it to the Designer. Client-requested features or changes outside the agreed scope remain subject to the written change-request process and do not become included work merely because client approval is required."
])

add_clause(doc, 9, "Intellectual property", paragraphs=[
    "Upon full payment, the Designer assigns to the Contractor all transferable intellectual-property rights in the final approved deliverables created specifically for the Project. The Contractor may transfer or license those rights to Leetlogic or the Contractor’s client.",
    "The Designer retains ownership of rejected concepts, unused drafts, general methods, know-how, and pre-existing tools. Any pre-existing or third-party asset included in a final deliverable remains subject to its applicable licence, which the Designer must disclose before use.",
    "The Designer may display the completed public work in a portfolio only after Leetlogic has publicly launched the relevant work or after obtaining the Contractor’s written permission. Confidential or unreleased materials must not be displayed."
])

add_clause(doc, 10, "Confidentiality", paragraphs=[
    "The Designer will keep confidential all non-public business, product, customer, technical, financial, and personal information received in connection with the Project. The Designer will use that information only to perform this Agreement, will share it only with authorized persons, and will take reasonable steps to prevent unauthorized access or disclosure.",
    "These obligations do not apply to information that is publicly available through no breach of this Agreement, was lawfully known to the Designer before disclosure, or must be disclosed by law. Where legally permitted, the Designer will notify the Contractor before a compelled disclosure."
])

add_clause(doc, 11, "Designer assurances", paragraphs=[
    "The Designer will perform the services professionally and with reasonable skill and care. The Designer represents that the final work will be original except for disclosed, properly licensed materials and that the Designer has authority to enter this Agreement and grant the rights stated here.",
    "The Designer will not knowingly introduce unlicensed assets, confidential materials belonging to another person, or design elements copied in a manner that infringes another party’s rights."
])

add_clause(doc, 12, "Independent contractor", paragraphs=[
    "The Designer is an independent contractor and not an employee, partner, agent, or joint venturer of the Contractor or Leetlogic. The Designer controls the manner and place of performing the services, subject to the agreed deliverables, standards, deadlines, and review process, and is responsible for the Designer’s own taxes and statutory obligations."
])

add_clause(doc, 13, "Termination", paragraphs=[
    "Either party may terminate this Agreement if the other party materially breaches it and does not remedy the breach within seven calendar days after written notice. The Contractor may also terminate the Project for convenience by written notice.",
    "If this Agreement is terminated before completion, the value of work earned will be determined by the deliverables completed and approved in writing by the Leetlogic client: ₦50,000 for Branding, ₦150,000 for Website Design, and ₦200,000 for App Design. A deliverable that has not been completed and approved by the client will not be treated as earned for this reconciliation unless the parties agree otherwise in writing.",
    "The parties will compare the approved value earned with all professional-fee payments already made. If the approved value exceeds the amount already paid, the Contractor will pay the difference. If payments already made exceed the approved value, the Designer will refund the difference. Any approved and documented ₦35,000 tools expense already incurred will be added separately to the amount due to the Designer and will not be reduced by the deliverable reconciliation.",
    "On termination, the Designer will promptly deliver all completed and paid-for work, editable files, and relevant project assets. Sections concerning confidentiality, intellectual property, payment obligations, and dispute resolution will survive termination."
])

add_clause(doc, 14, "Liability", paragraphs=[
    "Neither party will be liable to the other for indirect, special, or consequential loss arising from this Agreement. To the extent permitted by law, each party’s total liability connected with this Agreement will not exceed the total project fee, except for fraud, wilful misconduct, breach of confidentiality, infringement caused by unauthorized materials, or payment obligations."
])

add_clause(doc, 15, "Governing law and disputes", paragraphs=[
    "This Agreement is governed by the laws of the Federal Republic of Nigeria. The parties will first attempt in good faith to resolve any dispute through direct discussion. If unresolved within fourteen calendar days after written notice of the dispute, either party may pursue mediation or any other remedy available under Nigerian law."
])

add_clause(doc, 16, "General terms", paragraphs=[
    "This Agreement and the approved Leetlogic project brief contain the complete understanding between the parties about the services described here. If the project brief conflicts with this Agreement on fees, ownership, acceptance, confidentiality, or legal terms, this Agreement controls.",
    "Any amendment or waiver must be in writing and accepted by both parties. Email approval is sufficient for project decisions and change requests, but amendments to the total fee, ownership, or termination provisions should be signed by both parties.",
    "Neither party may assign this Agreement without the other party’s written consent, except that the Contractor may assign the final deliverables and related rights to Leetlogic or the Contractor’s client after full payment. Electronic signatures and counterparts are valid to the extent permitted by law."
])

doc.add_page_break()
h = doc.add_paragraph(style="Heading 1")
h.add_run("Signatures")
p = doc.add_paragraph("By signing below, each party confirms that they have read, understood, and agreed to this Agreement.")

sig = doc.add_table(rows=6, cols=2)
sig.alignment = WD_TABLE_ALIGNMENT.CENTER
sig.autofit = False
for c in sig.columns:
    c.width = Inches(3.25)
entries = [
    ("CONTRACTOR", "DESIGNER"),
    ("Great Ifeanyichukwu Adams", "David Chikwendu"),
    ("Signature: __________________________", "Signature: __________________________"),
    ("Date: ______________________________", "Date: ______________________________"),
    ("Email: kingifean@gmail.com", "Email: thedavelement@gmail.com"),
    ("Project: Leetlogic", "Project: Leetlogic"),
]
for i, row in enumerate(entries):
    for j, value in enumerate(row):
        set_cell_text(sig.cell(i, j), value, bold=(i in (0, 1)), size=10.5)
        if i == 0:
            set_cell_shading(sig.cell(i, j), "E8F1EC")

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("Leetlogic Design Services Agreement · Great Ifeanyichukwu Adams and David Chikwendu")
fr.font.name = "Aptos"
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor(90, 90, 90)

doc.core_properties.title = "Leetlogic Design Services Agreement"
doc.core_properties.subject = "Brand identity and UI/UX design services"
doc.core_properties.author = "Great Ifeanyichukwu Adams"
doc.save(OUTPUT)
print(OUTPUT)
